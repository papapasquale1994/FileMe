"""Crea i documenti di esempio usati per provare FileMe.

Tutti i nomi, gli indirizzi, le aziende e i numeri sono INVENTATI.
I file vengono scritti in esempi/documenti/, in sottocartelle che imitano un
computer reale (Download, Desktop, Documenti...). Alcuni nomi di file sono
volutamente poco chiari (es. "Allegato1.pdf"): è proprio il caso in cui la
ricerca semantica deve aiutare.

Uso (con l'ambiente virtuale attivo):  python esempi/genera_documenti.py
"""

from __future__ import annotations

import textwrap
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from docx import Document as DocumentoWord
from fpdf import FPDF
from openpyxl import Workbook

CARTELLA = Path(__file__).parent / "documenti"


@dataclass
class Documento:
    percorso: str  # dove salvarlo, relativo a esempi/documenti
    testo: str = ""  # righe che iniziano con "# " = titoli; "[tabella]" = posizione tabella (DOCX)
    tabella: list[list] | None = None  # solo DOCX
    fogli: dict[str, list[list]] | None = None  # solo XLSX: nome foglio -> righe
    codifica: str = "utf-8"  # solo TXT


def righe(testo: str) -> list[str]:
    """Toglie il rientro comune e restituisce il testo riga per riga."""
    return textwrap.dedent(testo).strip().splitlines()


def scrivi_pdf(doc: Documento, destinazione: Path) -> None:
    pdf = FPDF()
    pdf.set_creation_date(datetime(2024, 1, 1))
    pdf.add_page()
    for riga in righe(doc.testo):
        if riga.startswith("# "):
            pdf.set_font("Helvetica", style="B", size=13)
            pdf.multi_cell(0, 8, riga[2:], new_x="LMARGIN", new_y="NEXT")
        elif not riga.strip():
            pdf.ln(3)
        else:
            pdf.set_font("Helvetica", size=11)
            pdf.multi_cell(0, 6, riga, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(destinazione))


def scrivi_docx(doc: Documento, destinazione: Path) -> None:
    word = DocumentoWord()
    for riga in righe(doc.testo):
        if riga.startswith("# "):
            word.add_heading(riga[2:], level=1)
        elif riga.strip() == "[tabella]":
            tabella = word.add_table(rows=0, cols=len(doc.tabella[0]))
            tabella.style = "Table Grid"
            for valori in doc.tabella:
                for cella, valore in zip(tabella.add_row().cells, valori):
                    cella.text = str(valore)
        elif riga.strip():
            word.add_paragraph(riga)
    word.save(str(destinazione))


def scrivi_xlsx(doc: Documento, destinazione: Path) -> None:
    cartella_lavoro = Workbook()
    cartella_lavoro.remove(cartella_lavoro.active)  # tolgo il foglio vuoto iniziale
    for nome, righe_foglio in doc.fogli.items():
        foglio = cartella_lavoro.create_sheet(nome)
        for riga in righe_foglio:
            foglio.append(riga)
    cartella_lavoro.save(str(destinazione))


def scrivi_txt(doc: Documento, destinazione: Path) -> None:
    testo = "\n".join(righe(doc.testo)) + "\n"
    destinazione.write_text(testo, encoding=doc.codifica)


SCRITTORI = {".pdf": scrivi_pdf, ".docx": scrivi_docx, ".xlsx": scrivi_xlsx, ".txt": scrivi_txt}


# =============================================================================
# CONTENUTO DEI DOCUMENTI (tutto inventato)
# =============================================================================

DOCUMENTI = [
    # --- Download --------------------------------------------------------------
    Documento("Download/CV_2024_def.pdf", """
        # Giulia Ferri
        Via dei Mille 12, 40121 Bologna - Tel. 333 000 0000 - giulia.ferri@example.com
        Nata a Rimini il 14/05/1996 - Patente B, automunita

        # Profilo
        Receptionist con 5 anni di esperienza in hotel 4 stelle e strutture turistiche della Riviera romagnola e della Costiera amalfitana. Orientata all'accoglienza dell'ospite, abituata a lavorare su turni e in alta stagione. Cerco una posizione stabile nel settore turistico-alberghiero come front office o guest relations.

        # Esperienze professionali
        2022 - oggi: Receptionist, Hotel Villa Aurora (Rimini)
        Check-in e check-out, gestione prenotazioni con gestionale alberghiero e channel manager, assistenza agli ospiti stranieri, organizzazione di escursioni e transfer.

        2020 - 2021: Addetta all'accoglienza (stagionale), Residence Le Terrazze (Positano)
        Front desk, informazioni turistiche, gestione reclami, cassa e chiusura giornaliera.

        2018 - 2019: Animatrice turistica, villaggio vacanze in Calabria
        Animazione diurna per famiglie, mini club, spettacoli serali.

        # Formazione
        2018: Diploma di Istituto Tecnico per il Turismo, Rimini
        2023: Corso HACCP e corso di primo soccorso aziendale

        # Lingue
        Inglese C1 - Tedesco B2 - Spagnolo A2

        # Competenze
        Gestionali alberghieri, extranet dei portali di prenotazione, pacchetto Office, problem solving, lavoro in team.

        Autorizzo il trattamento dei dati personali ai sensi del Regolamento UE 2016/679 (GDPR).
    """),
    Documento("Download/Resume_Giulia_Ferri_EN.pdf", """
        # GIULIA FERRI
        Bologna, Italy - giulia.ferri@example.com - +39 333 000 0000

        # Summary
        Multilingual hospitality professional with five years of front office experience in four-star hotels and seasonal resorts. Looking for international opportunities with cruise lines or hotel chains abroad.

        # Experience
        2022 - present: Front Desk Agent, Hotel Villa Aurora, Rimini (Italy)
        - Managed check-in and check-out for up to 120 rooms per day
        - Handled bookings with the property management system and online travel agencies
        - Assisted international guests with tours and transfers

        2020 - 2021: Guest Services Assistant (seasonal), Le Terrazze Residence, Positano (Italy)

        # Education
        2018: High School Diploma in Tourism, Rimini

        # Languages
        Italian (native), English (C1), German (B2), Spanish (A2)

        # Certifications
        Food safety (HACCP), First Aid
    """),
    Documento("Download/fattura_luce_marzo_2024.pdf", """
        # Energia Chiara S.p.A. - Bolletta energia elettrica
        Cliente: Giulia Ferri - Codice cliente 00123456
        Indirizzo di fornitura: Via dei Mille 12, 40121 Bologna
        POD: IT001E00000000 - Offerta: Luce Casa Prezzo Fisso

        Periodo di fatturazione: 01/03/2024 - 31/03/2024
        Numero fattura: 2024/0331-778
        Data di emissione: 05/04/2024 - Scadenza: 25/04/2024

        # Riepilogo importi
        Spesa per la materia energia: 41,20 euro
        Spesa per il trasporto e la gestione del contatore: 12,85 euro
        Oneri di sistema: 3,10 euro
        Imposte e IVA 10%: 6,72 euro
        Totale da pagare: 63,87 euro

        # Consumi
        Consumo del periodo: 182 kWh (F1: 70 kWh, F2: 55 kWh, F3: 57 kWh)
        Consumo annuo stimato: 2.150 kWh

        Pagamento con addebito diretto sul conto corrente (SDD).
        Per segnalare guasti chiama il numero verde 800 000 000.
    """),
    Documento("Download/documento(3).pdf", """
        # GasNord Servizi - Fattura gas naturale
        Intestatario: Giulia Ferri
        Punto di riconsegna (PDR): 00000000000000 - Via dei Mille 12, Bologna
        Periodo: gennaio - febbraio 2024 (fattura bimestrale)

        Lettura precedente: 3.412 Smc - Lettura attuale: 3.598 Smc
        Consumo fatturato: 186 Smc (riscaldamento e acqua calda)

        Importo totale: 214,30 euro - Scadenza pagamento: 15/03/2024
        Bollettino pagabile in posta, in banca o con PagoPA.

        Consigli per risparmiare: abbassa il termostato di un grado e fai controllare la caldaia ogni anno.
    """),
    Documento("Download/stampa_20240312.pdf", """
        # Laboratorio Analisi Cliniche San Luca - Referto
        Paziente: Ferri Giulia - Data di nascita: 14/05/1996
        Data prelievo: 12/03/2024 - Medico richiedente: Dott.ssa Elena Neri

        # Emocromo
        Globuli rossi: 4,6 milioni/mm3 (valori di riferimento 4,2 - 5,4)
        Emoglobina: 13,1 g/dL (12,0 - 16,0)
        Globuli bianchi: 6.800 /mm3 (4.000 - 10.000)
        Piastrine: 245.000 /mm3 (150.000 - 400.000)

        # Chimica clinica
        Glicemia a digiuno: 88 mg/dL (70 - 100)
        Colesterolo totale: 212 mg/dL (desiderabile < 200) *
        Colesterolo HDL: 58 mg/dL
        Trigliceridi: 95 mg/dL (< 150)
        Ferritina: 18 ng/mL (15 - 150)

        * Valore fuori dall'intervallo di riferimento: si consiglia di discuterne con il proprio medico.
        Referto firmato digitalmente dal Direttore del laboratorio.
    """),
    Documento("Download/e-ticket_7F3KQ.pdf", """
        # SkyEuro Airlines - Booking confirmation and electronic ticket
        Booking reference: 7F3KQ
        Passenger: MS GIULIA FERRI

        # Outbound flight - 02 October 2024
        Bologna (BLQ) 10:40 -> Frankfurt (FRA) 12:15, flight SE 283
        Frankfurt (FRA) 13:30 -> Tokyo Haneda (HND) 08:25 +1 day, flight SE 716

        # Return flight - 13 October 2024
        Osaka Kansai (KIX) 10:15 -> Frankfurt (FRA) 16:40, flight SE 741
        Frankfurt (FRA) 18:05 -> Bologna (BLQ) 19:20, flight SE 290

        Fare: Economy Classic - 1 checked bag 23 kg, 1 hand luggage 8 kg
        Total paid: 1,148.60 EUR (credit card)

        Online check-in opens 23 hours before departure. Please bring a valid passport.
    """),
    Documento("Download/WM-7400_manual_EN.pdf", """
        # WM-7400 Front-loading washing machine - User manual (extract)

        # Before first use
        Remove the transport bolts from the back of the appliance before connecting it. Keep them for future transport. Connect the inlet hose to a cold water tap and make sure the drain hose is not bent.

        # Programs
        Cotton 40-90 °C: for towels and bed linen.
        Synthetics 40 °C: for shirts and mixed fabrics.
        Wool / Hand wash 30 °C: gentle drum movement, low spin.
        Quick 15 min: for lightly soiled clothes, max 2 kg.
        Eco 40-60: lowest energy consumption.

        # Cleaning the pump filter
        Clean the filter every three months or when the machine does not drain. Unplug the appliance, open the small door at the bottom right, place a low container under it and slowly unscrew the filter counter-clockwise. Remove fluff and coins, rinse the filter under running water and screw it back in firmly.

        # Troubleshooting
        Error E21: the water is not draining. Check the filter and the drain hose.
        Error E10: no water inlet. Check that the tap is open.
        The door does not open: wait two minutes after the end of the cycle.

        # Warranty
        The appliance is covered by a 2-year warranty from the date of purchase. Keep the receipt.
    """),
    # --- Desktop ---------------------------------------------------------------
    Documento("Desktop/Allegato1.pdf", """
        # CONTRATTO DI LOCAZIONE AD USO ABITATIVO
        (ai sensi dell'art. 2, comma 1, della Legge 9 dicembre 1998 n. 431)

        Il Sig. Paolo Bianchi, nato a Bologna il 03/02/1960 (di seguito "locatore"), concede in locazione alla Sig.ra Giulia Ferri, nata a Rimini il 14/05/1996 (di seguito "conduttrice"), che accetta, l'unità immobiliare sita in Bologna, Via dei Mille 12, piano 3, interno 7, composta da due camere, cucina abitabile, bagno e balcone, con cantina di pertinenza.

        Art. 1 - Durata. Il contratto è stipulato per la durata di quattro anni, dal 01/09/2023 al 31/08/2027, e si rinnova automaticamente per altri quattro anni salvo disdetta del locatore nei casi previsti dalla legge.

        Art. 2 - Canone. Il canone annuo è convenuto in euro 9.600,00, da pagarsi in 12 rate mensili anticipate di euro 800,00 ciascuna, entro il giorno 5 di ogni mese, tramite bonifico bancario.

        Art. 3 - Deposito cauzionale. A garanzia degli obblighi assunti la conduttrice versa un deposito cauzionale di euro 2.400,00, pari a tre mensilità, che sarà restituito al termine della locazione previa verifica dello stato dell'immobile.

        Art. 4 - Spese. Sono a carico della conduttrice le spese condominiali ordinarie, le utenze di luce, gas e acqua e la tassa sui rifiuti.

        Art. 5 - Recesso della conduttrice. La conduttrice può recedere in qualsiasi momento dal contratto dandone avviso al locatore con lettera raccomandata almeno sei mesi prima.

        Art. 6 - Divieto di sublocazione. È vietata la sublocazione totale o parziale dell'immobile senza il consenso scritto del locatore.

        Art. 7 - Registrazione. Il contratto viene registrato presso l'Agenzia delle Entrate con opzione per la cedolare secca.

        Letto, confermato e sottoscritto. Bologna, 25/08/2023.
        Il locatore ______________     La conduttrice ______________
    """),
    # Salvato con la codifica "Windows" (cp1252) come fa il vecchio Blocco note:
    # serve a verificare che leggiamo bene anche questi file.
    Documento("Desktop/note.txt", """
        TRASLOCO - cose da fare

        - disdire il contratto di luce e gas della vecchia casa (chiamare entro fine mese)
        - comunicare il cambio di residenza al Comune (anagrafe online)
        - aggiornare l'indirizzo su patente e libretto dell'auto
        - prenotare il furgone per sabato 26 agosto
        - chiedere a Luca e Sara se possono aiutare con i mobili
        - comprare scatole e nastro adesivo
        - fare le foto ai contatori il giorno della consegna delle chiavi
        - avvisare la banca e il medico di base del nuovo indirizzo
        - il citofono è rotto: lasciare il numero di cellulare al corriere
        - restituire le chiavi al vecchio proprietario e farsi ridare la caparra
    """, codifica="cp1252"),
    Documento("Desktop/ricetta_nonna.txt", """
        Ragù alla bolognese della nonna Rina (per 6 persone)

        Ingredienti:
        300 g di carne di manzo macinata
        200 g di pancetta tesa
        1 carota, 1 costa di sedano, mezza cipolla
        300 g di passata di pomodoro
        mezzo bicchiere di vino bianco secco
        un bicchiere di latte intero
        brodo, olio, sale e pepe

        Preparazione:
        Tritare finemente le verdure e la pancetta e farle rosolare in una pentola con poco olio.
        Aggiungere la carne e lasciarla rosolare bene, mescolando, finché non sfrigola.
        Sfumare con il vino e lasciarlo evaporare.
        Unire la passata, coprire e cuocere a fuoco bassissimo per almeno tre ore, aggiungendo brodo quando serve.
        Verso la fine versare il latte poco alla volta: la nonna diceva che è il segreto per un ragù morbido.
        Servire con tagliatelle fresche all'uovo, mai con gli spaghetti!
    """),
    Documento("Desktop/regali_natale.txt", """
        Idee regali Natale 2024

        Mamma: sciarpa di cashmere o un libro di ricette regionali
        Papà: abbonamento a una rivista di pesca, nuovo mulinello
        Marco: cuffie wireless (massimo 80 euro)
        Sara e Luca: cesto di prodotti tipici romagnoli
        Nonna Rina: cornice digitale con le foto di famiglia
        Colleghi della reception: panettone artigianale da dividere

        Budget totale: circa 400 euro. Comprare entro il 15 dicembre per le spedizioni.
    """),
    # --- Documenti/lavoro ------------------------------------------------------
    Documento("Documenti/lavoro/curriculum_vecchio.docx", """
        # Curriculum vitae - Giulia Ferri
        Bologna - giulia.ferri@example.com - 333 000 0000
        # Obiettivo
        Impiegata amministrativa con esperienza in contabilità di base, fatturazione e segreteria. Disponibile per contratti full-time in uffici amministrativi a Bologna e provincia.
        # Esperienze
        2019 - 2020: Impiegata amministrativa, Studio Commercialista Rossi & Associati (Bologna)
        Registrazione di fatture attive e passive, prima nota, archiviazione documenti, gestione dell'agenda clienti.
        2018: Stage in segreteria, Cooperativa Servizi Emilia
        Data entry, centralino, gestione della posta e del protocollo.
        # Formazione
        Diploma di Istituto Tecnico (2018). Corso di contabilità e software gestionale (2019).
        # Competenze
        [tabella]
    """, tabella=[
        ["Competenza", "Livello"],
        ["Excel", "avanzato"],
        ["Software di contabilità", "buono"],
        ["Fatturazione elettronica", "buono"],
        ["Inglese", "B2"],
    ]),
    Documento("Documenti/lavoro/lettera_presentazione_hotel.docx", """
        Giulia Ferri
        Via dei Mille 12, 40121 Bologna
        Spett.le Direzione del Personale
        Grand Hotel Lungomare - Riccione
        Oggetto: candidatura per la posizione di Front Office Manager
        Gentile Direttore,
        desidero propormi per la posizione di Front Office Manager pubblicata sul vostro sito. Da cinque anni lavoro alla reception di hotel quattro stelle e ho maturato esperienza nella gestione delle prenotazioni, nel coordinamento dei turni e nell'accoglienza di clientela internazionale.
        Parlo fluentemente inglese e tedesco e conosco i principali gestionali alberghieri. Mi piacerebbe mettere al servizio del vostro hotel la mia attenzione per l'ospite e la capacità di risolvere i problemi con calma anche nei momenti di alta stagione.
        Allego il mio curriculum e resto a disposizione per un colloquio, anche online.
        Cordiali saluti,
        Giulia Ferri
    """),
    Documento("Documenti/lavoro/reference_letter.docx", """
        Hotel Villa Aurora - Rimini, Italy
        To whom it may concern,
        I am pleased to recommend Ms Giulia Ferri, who has worked as a receptionist at our hotel since May 2022. During this time Giulia has proven to be reliable, punctual and extremely kind with our guests. She speaks excellent English and German and often helped colleagues with difficult bookings.
        Giulia trained two new seasonal employees and contributed to improving our online reviews. I am confident she would be an asset to any hospitality team.
        Please feel free to contact me for further information.
        Kind regards,
        Roberto Galli
        General Manager
    """),
    Documento("Documenti/lavoro/busta_paga_gennaio_2024.pdf", """
        # Cedolino paga - Gennaio 2024
        Azienda: Villa Aurora Hotel S.r.l. - Rimini
        Dipendente: Ferri Giulia - Matricola 0042
        Qualifica: Receptionist - Livello 4 - CCNL Turismo
        Contratto: tempo indeterminato, full-time 40 ore settimanali

        # Competenze
        Paga base: 1.540,00 euro
        Contingenza: 524,00 euro
        Maggiorazioni lavoro festivo e notturno: 96,40 euro
        Totale competenze lorde: 2.160,40 euro

        # Trattenute
        Contributi INPS a carico del dipendente: 198,50 euro
        IRPEF netta: 285,30 euro
        Addizionali regionale e comunale: 41,10 euro
        Totale trattenute: 524,90 euro

        NETTO IN BUSTA: 1.635,50 euro
        Ferie residue: 12 giorni - Permessi residui: 20 ore
        TFR maturato nel mese: 160,03 euro
    """),
    Documento("Documenti/lavoro/turni_reception_luglio.xlsx", fogli={
        "Turni luglio 2024": [
            ["Data", "Mattina 7-15", "Pomeriggio 15-23", "Notte 23-7"],
            [date(2024, 7, 1), "Giulia", "Marta", "Davide"],
            [date(2024, 7, 2), "Giulia", "Marta", "Davide"],
            [date(2024, 7, 3), "Marta", "Kevin", "Davide"],
            [date(2024, 7, 4), "Marta", "Giulia", "Kevin"],
            [date(2024, 7, 5), "Kevin", "Giulia", "Davide"],
            [date(2024, 7, 6), "Kevin", "Giulia", "Marta"],
            [date(2024, 7, 7), "Davide", "Kevin", "Marta"],
        ],
        "Note": [
            ["Nota"],
            ["Ferie di Giulia dal 22 al 26 luglio"],
            ["Cambi turno: chiedere al responsabile con 48 ore di anticipo"],
            ["Durante il turno di notte fare la chiusura di cassa entro le 3"],
        ],
    }),
    # --- Documenti/casa --------------------------------------------------------
    Documento("Documenti/casa/spese_condominio_2024.xlsx", fogli={
        "Riparto 2024": [
            ["Condominio Via dei Mille 12 - Rendiconto spese ordinarie 2024"],
            [],
            ["Voce di spesa", "Importo totale (euro)", "Millesimi interno 7", "Quota interno 7 (euro)"],
            ["Pulizia scale", 2400, 42, 100.80],
            ["Luce parti comuni", 650, 42, 27.30],
            ["Manutenzione ascensore", 1800, 42, 75.60],
            ["Compenso amministratore", 1500, 42, 63.00],
            ["Assicurazione fabbricato", 900, 42, 37.80],
            ["Acqua", 3200, 42, 134.40],
            ["Totale", 10450, 42, 438.90],
        ],
    }),
    Documento("Documenti/casa/budget_famiglia_2024.xlsx", fogli={
        "Spese mensili": [
            ["Categoria", "Gennaio", "Febbraio", "Marzo", "Totale trimestre"],
            ["Affitto", 800, 800, 800, 2400],
            ["Spesa alimentare", 420, 390, 450, 1260],
            ["Luce", 58, 61, 64, 183],
            ["Gas", 110, 104, 72, 286],
            ["Internet e telefono", 35, 35, 35, 105],
            ["Trasporti e benzina", 60, 75, 60, 195],
            ["Palestra", 39, 39, 39, 117],
            ["Svago e ristoranti", 120, 90, 150, 360],
            ["Totale", 1642, 1594, 1670, 4906],
        ],
        "Entrate": [
            ["Voce", "Importo mensile netto"],
            ["Stipendio Giulia", 1635.50],
            ["Stipendio Marco", 1720.00],
        ],
    }),
    Documento("Documenti/casa/inventario_trasloco.xlsx", fogli={
        "Inventario": [
            ["Scatola o oggetto", "Stanza di destinazione", "Contenuto", "Fragile"],
            ["Scatola 1", "Cucina", "Piatti e bicchieri", "Sì"],
            ["Scatola 2", "Cucina", "Pentole e padelle", "No"],
            ["Scatola 3", "Camera", "Vestiti invernali", "No"],
            ["Scatola 4", "Studio", "Libri e documenti", "No"],
            ["Scatola 5", "Bagno", "Asciugamani e prodotti per la casa", "No"],
            ["Divano letto", "Soggiorno", "", "No"],
            ["Specchio grande", "Ingresso", "", "Sì"],
            ["Televisore 43 pollici", "Soggiorno", "", "Sì"],
            ["Lavatrice", "Bagno", "Ricordarsi i bulloni per il trasporto", "No"],
        ],
    }),
    # --- Documenti/auto, tasse, viaggi, studio ---------------------------------
    Documento("Documenti/auto/polizza_rc_auto_2024.pdf", """
        # Sicura Mutua Assicurazioni - Polizza RC Auto
        Numero polizza: RCA-2024-558812
        Contraente e proprietaria: Giulia Ferri
        Veicolo: utilitaria 1.2 benzina - Targa: AB000CD
        Data di immatricolazione: 06/2017

        Periodo di copertura: dalle ore 24 del 14/02/2024 alle ore 24 del 14/02/2025
        Premio annuo: 412,00 euro (rata unica)
        Classe di merito CU: 4

        # Massimali
        Danni alle persone: 6.450.000 euro per sinistro
        Danni alle cose: 1.300.000 euro per sinistro

        # Garanzie accessorie
        Assistenza stradale 24 ore su 24 con carro attrezzi
        Tutela legale
        Infortuni del conducente

        In caso di incidente compila il modulo di constatazione amichevole (CAI) e comunicalo entro 3 giorni.
    """),
    Documento("Documenti/tasse/promemoria_730_2024.docx", """
        # Promemoria dichiarazione dei redditi 2024 (modello 730, redditi 2023)
        Appuntamento al CAF: martedì 14 maggio alle 10:30. Portare documento d'identità e tessera sanitaria.
        Documenti da portare:
        - Certificazione Unica (CU) 2024 del datore di lavoro
        - contratto di affitto registrato (per la detrazione per giovani inquilini)
        - ricevute delle spese mediche e scontrini parlanti della farmacia
        - le ricevute della palestra non servono (non sono detraibili)
        Nota: chiedere se si possono detrarre anche gli occhiali da vista.
        # Spese mediche 2023 (detrazione del 19% oltre la franchigia di 129,11 euro)
        [tabella]
    """, tabella=[
        ["Data", "Descrizione", "Importo (euro)"],
        ["15/02/2023", "Visita oculistica", "120,00"],
        ["03/04/2023", "Farmaci (scontrini della farmacia)", "86,40"],
        ["20/06/2023", "Analisi del sangue", "54,00"],
        ["11/10/2023", "Fisioterapia (5 sedute)", "250,00"],
        ["", "Totale", "510,40"],
    ]),
    Documento("Documenti/viaggi/itinerario_giappone.docx", """
        # Viaggio in Giappone - dal 2 al 13 ottobre 2024
        # Giorni 1-4: Tokyo
        Arrivo all'aeroporto di Haneda. Hotel a Shinjuku. Visita a Shibuya, Asakusa e al tempio Senso-ji, mercato esterno di Tsukiji, quartiere di Akihabara. Gita di un giorno a Nikko.
        # Giorno 5: Hakone
        Treno per Hakone, giro del lago Ashi in battello, vista sul monte Fuji se il tempo è bello. Notte in ryokan con onsen.
        # Giorni 6-9: Kyoto
        Santuario Fushimi Inari all'alba per evitare la folla, foresta di bambù di Arashiyama, Padiglione d'oro (Kinkaku-ji), quartiere di Gion la sera. Gita a Nara per i cervi e il Grande Buddha.
        # Giorni 10-11: Osaka
        Street food a Dotonbori, castello di Osaka. Partenza dall'aeroporto del Kansai.
        # Da ricordare
        Japan Rail Pass da 7 giorni da attivare a Tokyo. Adattatore per prese di tipo A. Portare contanti: molti locali non accettano carte. Scaricare le mappe offline.
    """),
    Documento("Documenti/studio/attestato_haccp.pdf", """
        # ATTESTATO DI FREQUENZA
        Corso di formazione per addetti alla manipolazione degli alimenti (HACCP)
        ai sensi del Regolamento CE 852/2004

        Si attesta che GIULIA FERRI, nata a Rimini il 14/05/1996, ha frequentato il corso di formazione in materia di igiene e sicurezza alimentare della durata di 12 ore, superando la verifica finale di apprendimento.

        Argomenti trattati: rischi di contaminazione degli alimenti, corretta conservazione e catena del freddo, pulizia e sanificazione, allergeni, autocontrollo e piano HACCP.

        Validità: 2 anni. Rinnovo consigliato entro marzo 2025.
        Ente di formazione: Centro Formazione Turismo Romagna - Rimini, 18/03/2023
    """),
    # Documento lungo: parla molto di turismo ma NON è un curriculum.
    # Serve a verificare che la ricerca non si faccia ingannare dalle parole.
    Documento("Documenti/studio/tesi_cap2_turismo_sostenibile.docx", """
        # Capitolo 2 - Il turismo sostenibile nelle destinazioni costiere
        # 2.1 Definizioni e origini del concetto
        Il concetto di turismo sostenibile nasce alla fine degli anni Ottanta, sulla scia del più ampio dibattito sullo sviluppo sostenibile. Secondo la definizione più diffusa, è sostenibile il turismo che tiene pienamente conto dei suoi impatti economici, sociali e ambientali, attuali e futuri, rispondendo ai bisogni dei visitatori, del settore, dell'ambiente e delle comunità ospitanti.
        Questa definizione mette in evidenza tre dimensioni. La dimensione ambientale riguarda l'uso responsabile delle risorse naturali, come l'acqua, il suolo e il paesaggio. La dimensione sociale riguarda il rispetto dell'identità culturale delle comunità locali e la qualità della vita dei residenti. La dimensione economica, infine, richiede che il turismo generi benefici duraturi e distribuiti in modo equo, compresa un'occupazione stabile e non soltanto stagionale.
        # 2.2 Il problema della stagionalità
        Le destinazioni costiere del Mediterraneo sono caratterizzate da una forte concentrazione dei flussi nei mesi estivi. In molte località della Riviera adriatica oltre il settanta per cento delle presenze annuali si registra tra giugno e agosto. Questa concentrazione produce effetti negativi su più fronti: sovraccarico delle infrastrutture idriche e dei trasporti, congestione delle spiagge, aumento dei rifiuti e pressione sui prezzi degli affitti per i residenti.
        Dal punto di vista del lavoro, la stagionalità si traduce in contratti brevi e in una elevata rotazione del personale nelle strutture ricettive. Molti lavoratori del settore alternano mesi di attività intensa a lunghi periodi di disoccupazione, con conseguenze sulla qualità del servizio e sulla possibilità di investire nella formazione.
        # 2.3 Indicatori per misurare la sostenibilità
        Per valutare la sostenibilità di una destinazione non bastano i dati sugli arrivi e sulle presenze. La letteratura propone sistemi di indicatori che combinano misure ambientali, come i consumi idrici ed energetici per presenza turistica, la quota di rifiuti avviati a riciclo e la qualità delle acque di balneazione, con misure sociali ed economiche, come il rapporto tra turisti e residenti, la durata media dei contratti di lavoro e la soddisfazione della popolazione locale.
        Il Sistema europeo di indicatori per il turismo (ETIS) è uno degli strumenti più utilizzati dalle amministrazioni locali. Il suo pregio principale è la semplicità: i comuni possono raccogliere i dati con risorse limitate e confrontarsi nel tempo con altre destinazioni simili.
        # 2.4 Buone pratiche: il caso della Riviera romagnola
        Negli ultimi anni alcuni comuni della Riviera romagnola hanno avviato strategie per destagionalizzare l'offerta, puntando sul turismo congressuale, sugli eventi sportivi e sul cicloturismo nell'entroterra. Diverse strutture alberghiere hanno ottenuto certificazioni ambientali, riducendo i consumi energetici con pannelli solari e sistemi di recupero dell'acqua piovana.
        Particolarmente interessante è la collaborazione tra albergatori e scuole professionali per la formazione del personale durante la bassa stagione: in questo modo i lavoratori possono aggiornare le proprie competenze linguistiche e digitali, e le imprese possono contare su collaboratori qualificati che tornano anno dopo anno.
        Questi esempi mostrano che la sostenibilità non è solo un vincolo, ma può diventare un vantaggio competitivo per le destinazioni che sanno innovare. Nel capitolo successivo si analizzeranno i risultati di un questionario somministrato a un campione di operatori turistici della provincia di Rimini.
    """),
]


def main() -> None:
    for doc in DOCUMENTI:
        destinazione = CARTELLA / doc.percorso
        destinazione.parent.mkdir(parents=True, exist_ok=True)
        SCRITTORI[destinazione.suffix](doc, destinazione)
        print(f"creato  {doc.percorso}")
    print(f"\n{len(DOCUMENTI)} documenti di esempio in {CARTELLA}")


if __name__ == "__main__":
    main()
