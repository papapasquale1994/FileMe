"""Estrazione: dai file sul disco al testo diviso in chunk.

Il percorso di ogni file è:
    trova_file()  ->  leggi il testo (PDF/DOCX/XLSX/TXT)  ->  pulisci()  ->  dividi_in_chunk()

Tutti i file vengono aperti in SOLA LETTURA (modalità "rb"): non vengono mai
modificati, spostati o cancellati.
"""

from __future__ import annotations

import codecs
import logging
import os
import re
from dataclasses import dataclass
from datetime import datetime, time
from pathlib import Path
from typing import BinaryIO, Iterator

from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from pypdf import PdfReader

from fileme import config

# pypdf stampa molti avvisi tecnici sui PDF un po' "storti": li nascondiamo.
logging.getLogger("pypdf").setLevel(logging.ERROR)


@dataclass
class DocumentoEstratto:
    """Il risultato dell'estrazione di un file."""

    percorso: Path
    tipo: str  # "pdf", "docx", "xlsx" o "txt"
    chunk: list[str]  # i pezzi di testo; vuota se il file non contiene testo
    errore: str | None = None  # descrizione del problema, se il file non si è potuto leggere


# --- 1. Trovare i file ---------------------------------------------------------


def trova_file(percorso: Path) -> list[Path]:
    """Elenca i file supportati in una cartella e nelle sue sottocartelle.

    Salta i file e le cartelle nascosti (che iniziano con ".") e i file
    temporanei di Office (che iniziano con "~$").
    Accetta anche il percorso di un singolo file.
    """
    if percorso.is_file():
        return [percorso] if _supportato(percorso.name) else []

    trovati = []
    for cartella, sottocartelle, file in os.walk(percorso):
        sottocartelle[:] = [c for c in sottocartelle if not c.startswith(".")]
        trovati += [Path(cartella) / nome for nome in file if _supportato(nome)]
    return sorted(trovati)


def _supportato(nome: str) -> bool:
    if nome.startswith((".", "~$")):
        return False
    return Path(nome).suffix.lower() in config.ESTENSIONI_SUPPORTATE


# --- 2. Leggere il testo, un formato alla volta ----------------------------------


def _leggi_pdf(file: BinaryIO) -> str:
    lettore = PdfReader(file)
    # Molti PDF sono "protetti" con una password vuota: proviamo ad aprirli.
    if lettore.is_encrypted and not lettore.decrypt(""):
        raise ValueError("PDF protetto da password")
    # Le pagine diventano paragrafi separati (riga vuota tra una e l'altra).
    return "\n\n".join(pagina.extract_text() or "" for pagina in lettore.pages)


def _leggi_docx(file: BinaryIO) -> str:
    documento = Document(file)
    parti = []
    # Paragrafi e tabelle nell'ordine in cui compaiono nel documento.
    for blocco in documento.iter_inner_content():
        if isinstance(blocco, Table):
            parti.append(_testo_tabella(blocco))
        else:
            parti.append(blocco.text)
    return "\n\n".join(parti)


def _testo_tabella(tabella: Table) -> str:
    """Una riga di testo per ogni riga della tabella: "cella | cella | cella"."""
    righe = []
    for riga in tabella.rows:
        celle: list[str] = []
        for cella in riga.cells:
            testo = cella.text.strip()
            # Le celle unite ripetono lo stesso testo: lo teniamo una volta sola.
            if testo and (not celle or celle[-1] != testo):
                celle.append(testo)
        righe.append(" | ".join(celle))
    return "\n".join(righe)


def _leggi_xlsx(file: BinaryIO) -> str:
    # data_only=True: per le formule leggiamo il risultato, non la formula.
    cartella_lavoro = load_workbook(file, read_only=True, data_only=True)
    parti = []
    try:
        for foglio in cartella_lavoro.worksheets:
            righe = [f"Foglio: {foglio.title}"]
            for valori in foglio.iter_rows(values_only=True):
                celle = [_testo_cella(v) for v in valori if v is not None and str(v).strip()]
                if celle:
                    righe.append(" | ".join(celle))
                if len(righe) > config.MAX_RIGHE_PER_FOGLIO:
                    break
            parti.append("\n".join(righe))
    finally:
        cartella_lavoro.close()
    return "\n\n".join(parti)


def _testo_cella(valore: object) -> str:
    """Scrive il valore di una cella in modo leggibile (es. date come 01/07/2024)."""
    if isinstance(valore, datetime):
        formato = "%d/%m/%Y" if valore.time() == time(0) else "%d/%m/%Y %H:%M"
        return valore.strftime(formato)
    if isinstance(valore, float) and valore.is_integer():
        return str(int(valore))
    return str(valore).strip()


def _leggi_txt(file: BinaryIO) -> str:
    dati = file.read()
    # I file di testo possono essere salvati con "codifiche" diverse.
    # Proviamo le più comuni: UTF-16 (riconoscibile dal suo marcatore iniziale),
    # UTF-8 (lo standard) e cp1252 (il vecchio standard di Windows).
    if dati.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return dati.decode("utf-16")
    for codifica in ("utf-8-sig", "cp1252"):
        try:
            return dati.decode(codifica)
        except UnicodeDecodeError:
            continue
    return dati.decode("latin-1")  # non fallisce mai: ultima spiaggia


LETTORI = {".pdf": _leggi_pdf, ".docx": _leggi_docx, ".xlsx": _leggi_xlsx, ".txt": _leggi_txt}


# --- 3. Pulire e dividere il testo ---------------------------------------------


def pulisci(testo: str) -> str:
    """Uniforma spazi e a capo, e toglie le righe vuote in eccesso."""
    testo = testo.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    righe = [" ".join(riga.split()) for riga in testo.split("\n")]  # spazi multipli -> uno
    testo = "\n".join(righe)
    testo = re.sub(r"\n{3,}", "\n\n", testo)  # al massimo una riga vuota di fila
    return testo.strip()


def dividi_in_chunk(
    testo: str,
    max_parole: int = config.PAROLE_PER_CHUNK,
    sovrapposizione: int = config.PAROLE_SOVRAPPOSIZIONE,
) -> list[str]:
    """Divide il testo in chunk di al massimo `max_parole` parole.

    Cerca di tagliare tra un paragrafo e l'altro. Un paragrafo troppo lungo
    viene spezzato. Ogni chunk ripete le ultime `sovrapposizione` parole del
    precedente, per non perdere il contesto delle frasi tagliate.
    """
    paragrafi = [p for p in re.split(r"\n\s*\n", testo) if p.strip()]

    # Spezzo i paragrafi lunghi in pezzi che, sommati alla sovrapposizione,
    # stanno comunque dentro il limite.
    lunghezza_pezzo = max_parole - sovrapposizione
    pezzi = []
    for paragrafo in paragrafi:
        parole = paragrafo.split()
        for inizio in range(0, len(parole), lunghezza_pezzo):
            pezzi.append(parole[inizio : inizio + lunghezza_pezzo])

    chunk = []
    corrente: list[str] = []
    for pezzo in pezzi:
        if corrente and len(corrente) + len(pezzo) > max_parole:
            chunk.append(" ".join(corrente))
            corrente = corrente[-sovrapposizione:]
        corrente += pezzo
    if corrente:
        chunk.append(" ".join(corrente))
    return chunk


# --- Mettiamo tutto insieme ----------------------------------------------------


def estrai_documento(percorso: Path) -> DocumentoEstratto:
    """Legge un file e restituisce il suo testo diviso in chunk.

    Se il file è danneggiato o illeggibile non si blocca: restituisce un
    DocumentoEstratto con la descrizione dell'errore.
    """
    estensione = percorso.suffix.lower()
    tipo = estensione.lstrip(".")
    try:
        with open(percorso, "rb") as file:  # "rb" = solo lettura
            testo = LETTORI[estensione](file)
    except Exception as errore:  # qualsiasi problema su un file non deve fermare gli altri
        return DocumentoEstratto(percorso, tipo, [], errore=f"{type(errore).__name__}: {errore}")
    return DocumentoEstratto(percorso, tipo, dividi_in_chunk(pulisci(testo)))


def estrai_cartella(percorso: Path) -> Iterator[DocumentoEstratto]:
    """Estrae, uno alla volta, tutti i file supportati di una cartella."""
    for file in trova_file(percorso):
        yield estrai_documento(file)
