"""Test dell'estrazione: ogni formato si legge, il testo si divide bene in chunk,
i file danneggiati non bloccano nulla e i file originali restano intatti."""

import hashlib
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from fpdf import FPDF
from openpyxl import Workbook
from pypdf import PdfWriter

from fileme.cli import main
from fileme.estrazione import (
    dividi_in_chunk,
    estrai_cartella,
    estrai_documento,
    pulisci,
    trova_file,
)

CARTELLA_ESEMPI = Path(__file__).parent.parent / "esempi" / "documenti"


# --- Piccoli aiutanti che creano file di prova in una cartella temporanea -------


def crea_pdf(percorso: Path, righe: list[str]) -> Path:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for riga in righe:
        pdf.multi_cell(0, 6, riga, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(percorso))
    return percorso


def crea_docx(percorso: Path) -> Path:
    word = Document()
    word.add_heading("Curriculum vitae", level=1)
    word.add_paragraph("Esperienza come receptionist in hotel.")
    tabella = word.add_table(rows=2, cols=2)
    tabella.cell(0, 0).text = "Lingua"
    tabella.cell(0, 1).text = "Livello"
    tabella.cell(1, 0).text = "Inglese"
    tabella.cell(1, 1).text = "C1"
    word.add_paragraph("Disponibile da subito.")
    word.save(str(percorso))
    return percorso


def crea_xlsx(percorso: Path) -> Path:
    cartella_lavoro = Workbook()
    foglio = cartella_lavoro.active
    foglio.title = "Turni"
    foglio.append(["Data", "Persona"])
    foglio.append([date(2024, 7, 1), "Giulia"])
    foglio.append([None, None])  # riga vuota: va saltata
    cartella_lavoro.create_sheet("Note").append(["Ferie in agosto", 3.0])
    cartella_lavoro.save(str(percorso))
    return percorso


def impronta(percorso: Path) -> tuple[str, float]:
    """Contenuto (hash) e data di modifica: se cambiano, il file è stato toccato."""
    return hashlib.sha256(percorso.read_bytes()).hexdigest(), percorso.stat().st_mtime


# --- Lettura dei singoli formati ---------------------------------------------------


def test_pdf(tmp_path):
    pdf = crea_pdf(tmp_path / "bolletta.pdf", ["Bolletta della luce", "Totale da pagare: 63,87 euro"])
    doc = estrai_documento(pdf)
    assert doc.errore is None
    assert doc.tipo == "pdf"
    assert "Bolletta della luce" in doc.chunk[0]
    assert "63,87 euro" in doc.chunk[0]


def cifra_pdf(origine: Path, destinazione: Path, password_apertura: str) -> Path:
    scrittore = PdfWriter(clone_from=str(origine))
    scrittore.encrypt(user_password=password_apertura, owner_password="proprietario", algorithm="AES-256")
    scrittore.write(str(destinazione))
    return destinazione


def test_pdf_protetto_senza_password_si_legge(tmp_path):
    # Come molti estratti conto: cifrato, ma si apre senza chiedere password
    originale = crea_pdf(tmp_path / "originale.pdf", ["Estratto conto bancario"])
    doc = estrai_documento(cifra_pdf(originale, tmp_path / "protetto.pdf", password_apertura=""))
    assert doc.errore is None
    assert "Estratto conto bancario" in doc.chunk[0]


def test_pdf_con_password_da_errore_chiaro(tmp_path):
    originale = crea_pdf(tmp_path / "originale.pdf", ["Segreto"])
    doc = estrai_documento(cifra_pdf(originale, tmp_path / "segreto.pdf", password_apertura="1234"))
    assert doc.chunk == []
    assert "password" in doc.errore


def test_pdf_senza_testo_non_e_un_errore(tmp_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.rect(10, 10, 50, 50)  # solo un disegno, come una scansione
    pdf.output(str(tmp_path / "scansione.pdf"))
    doc = estrai_documento(tmp_path / "scansione.pdf")
    assert doc.errore is None
    assert doc.chunk == []


def test_docx_con_tabella_in_ordine(tmp_path):
    doc = estrai_documento(crea_docx(tmp_path / "cv.docx"))
    testo = doc.chunk[0]
    assert "Lingua | Livello" in testo
    assert "Inglese | C1" in testo
    # La tabella resta tra il paragrafo che la precede e quello che la segue
    assert testo.index("receptionist") < testo.index("Inglese") < testo.index("Disponibile")


def test_xlsx_fogli_date_e_numeri(tmp_path):
    doc = estrai_documento(crea_xlsx(tmp_path / "turni.xlsx"))
    testo = doc.chunk[0]
    assert "Foglio: Turni" in testo
    assert "Foglio: Note" in testo
    assert "01/07/2024 | Giulia" in testo  # data in formato italiano
    assert "Ferie in agosto | 3" in testo  # 3.0 diventa 3
    assert "None" not in testo


@pytest.mark.parametrize("codifica", ["utf-8", "utf-8-sig", "cp1252", "utf-16"])
def test_txt_con_codifiche_diverse(tmp_path, codifica):
    percorso = tmp_path / "nota.txt"
    percorso.write_bytes("Il citofono è rotto, però funziona.".encode(codifica))
    doc = estrai_documento(percorso)
    assert doc.chunk == ["Il citofono è rotto, però funziona."]


def test_file_danneggiato_non_blocca(tmp_path):
    rotto = tmp_path / "rotto.docx"
    rotto.write_bytes(b"questo non e' un vero file Word")
    doc = estrai_documento(rotto)
    assert doc.chunk == []
    assert doc.errore  # c'è una descrizione del problema


# --- Ricerca dei file nelle cartelle ---------------------------------------------


def test_trova_file_supportati_e_salta_quelli_da_ignorare(tmp_path):
    (tmp_path / "sotto" / "cartella").mkdir(parents=True)
    (tmp_path / ".nascosta").mkdir()
    for nome in [
        "a.pdf",
        "B.PDF",  # estensione maiuscola: va presa
        "sotto/cartella/c.txt",
        "foto.jpg",  # formato non supportato
        "~$documento.docx",  # file temporaneo di Word
        ".segreto.txt",  # file nascosto
        ".nascosta/d.txt",  # dentro una cartella nascosta
    ]:
        (tmp_path / nome).write_text("x")
    trovati = {p.relative_to(tmp_path).as_posix() for p in trova_file(tmp_path)}
    assert trovati == {"a.pdf", "B.PDF", "sotto/cartella/c.txt"}


def test_trova_file_accetta_un_singolo_file(tmp_path):
    (tmp_path / "nota.txt").write_text("ciao")
    assert trova_file(tmp_path / "nota.txt") == [tmp_path / "nota.txt"]


# --- Pulizia e divisione in chunk ------------------------------------------------


def test_pulisci_spazi_e_righe_vuote():
    assert pulisci("  ciao   mondo \r\n\r\n\r\n\r\nfine  ") == "ciao mondo\n\nfine"


def test_testo_corto_un_solo_chunk():
    assert dividi_in_chunk("Una frase breve.") == ["Una frase breve."]
    assert dividi_in_chunk("") == []


def test_testo_lungo_diviso_con_sovrapposizione():
    parole = [f"p{i}" for i in range(600)]
    chunk = dividi_in_chunk(" ".join(parole), max_parole=250, sovrapposizione=50)
    assert len(chunk) > 1
    for testo in chunk:
        assert len(testo.split()) <= 250
    # Le ultime 50 parole di un chunk aprono il successivo
    for precedente, successivo in zip(chunk, chunk[1:]):
        assert successivo.split()[:50] == precedente.split()[-50:]
    # Nessuna parola persa
    assert chunk[-1].split()[-1] == "p599"


def test_taglia_tra_paragrafi():
    primo = " ".join(["uno"] * 150)
    secondo = " ".join(["due"] * 150)
    chunk = dividi_in_chunk(f"{primo}\n\n{secondo}", max_parole=250, sovrapposizione=50)
    assert chunk[0] == primo  # il primo paragrafo non viene spezzato a metà


# --- Sicurezza e documenti di esempio --------------------------------------------


def test_i_file_non_vengono_modificati(tmp_path):
    file = [
        crea_pdf(tmp_path / "a.pdf", ["prova"]),
        crea_docx(tmp_path / "b.docx"),
        crea_xlsx(tmp_path / "c.xlsx"),
    ]
    prima = {f: impronta(f) for f in file}
    list(estrai_cartella(tmp_path))
    assert {f: impronta(f) for f in file} == prima
    assert sorted(tmp_path.iterdir()) == sorted(file)  # nessun file nuovo creato


def test_tutti_i_documenti_di_esempio_si_leggono():
    documenti = list(estrai_cartella(CARTELLA_ESEMPI))
    assert len(documenti) == 24
    for doc in documenti:
        assert doc.errore is None, f"{doc.percorso}: {doc.errore}"
        assert doc.chunk, f"{doc.percorso}: nessun testo"


# --- Comando da terminale ------------------------------------------------------


def test_comando_estrai(tmp_path, capsys):
    crea_pdf(tmp_path / "bolletta.pdf", ["Bolletta della luce"])
    (tmp_path / "rotto.docx").write_bytes(b"non sono un docx")
    main(["estrai", str(tmp_path)])
    uscita = capsys.readouterr().out
    assert "[PDF ] bolletta.pdf  ->  1 chunk" in uscita
    assert "Bolletta della luce" in uscita
    assert "ERRORE" in uscita
    assert "file letti: 2 | chunk: 1 | senza testo: 0 | errori: 1" in uscita


def test_comando_estrai_cartella_inesistente(tmp_path):
    with pytest.raises(SystemExit):
        main(["estrai", str(tmp_path / "non_esiste")])
